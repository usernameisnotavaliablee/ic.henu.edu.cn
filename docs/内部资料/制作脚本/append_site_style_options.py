"""Append one editable style-selection page without rebuilding the reviewed CVs."""
from pathlib import Path
import sys
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
docs = next(p for p in Path(__file__).resolve().parents if p.name=='docs')
source = docs / '历史版本' / '三位教师简历与照片_人工核验稿.docx'
output = docs / '历史版本' / '三位教师简历与照片_含网站风格选型.docx'
doc = Document(source)
original_paragraph_count = len(doc.paragraphs)
original_blocks = [etree.tostring(e) for e in doc._element.body if e.tag != qn('w:sectPr')]

heading = doc.add_heading('学院官网设计风格选型', level=1)
heading.paragraph_format.page_break_before = True
p = doc.add_paragraph('请从以下四种方向中选择偏好的主风格，也可注明希望组合的特点。具体颜色以学校或学院视觉规范为准，栏目内容按后续确认清单安排。')

options = [
    ('A  稳重学院风', ['8B2835', 'ECE8E3', 'FFFFFF'], '深红、白、浅灰',
     '正式、稳重，接近传统高校官网。',
     '校院标识突出；横幅实景照片；新闻与通知分栏排列，信息集中。',
     '校院标识及少量实景照片。\n相对工作量：较低。'),
    ('B  现代科技风', ['173F67', '4474C2', 'E7F1F8'], '藏蓝、科技蓝、白',
     '清晰、现代，突出集成电路学科特色。',
     '芯片或实验室主题首图；浅色正文；研究方向、师资和新闻分区展示。',
     '一张可用的芯片主题图或实验室照片。\n静态版本工作量：中等。'),
    ('C  简洁学术风', ['FFFFFF', 'E7EAEE', '315779'], '白、浅灰、少量蓝',
     '简洁、易读，以学院内容为重点。',
     '较小首图；清楚的标题和留白；教师、课程、科研内容便于浏览。',
     '完整文字、教师照片及少量配图。\n相对工作量：较低。'),
    ('D  大图展示风', ['102633', 'B4C8D2', 'FFFFFF'], '深色首屏、浅色正文',
     '视觉突出，适合展示学院形象。',
     '大幅照片配简短介绍；下方排列内容板块。视频可作为后续选项。',
     '高质量横幅大图；选视频时需提供可用视频。\n相对工作量：中等至较高。'),
]

t = doc.add_table(rows=1, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [1.35, 1.22, 2.2, 2.03]
headers = ['风格与配色', '整体感觉', '首页大致效果', '素材和工作量']
for col, width in zip(t.columns, widths): col.width = Inches(width)
pr = t._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
    b = OxmlElement('w:'+side)
    for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): b.set(qn('w:'+k),v)
    borders.append(b)
pr.append(borders)
margin = OxmlElement('w:tblCellMar')
for side, val in [('top',120),('bottom',120),('left',105),('right',105)]:
    node=OxmlElement('w:'+side);node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa');margin.append(node)
pr.append(margin)

def cell_text(cell, text, bold=False):
    p=cell.paragraphs[0]
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1.12
    r=p.add_run(text);r.font.size=Pt(10.5);r.bold=bold
    return p

for i, (cell, label) in enumerate(zip(t.rows[0].cells, headers)):
    cell.width=Inches(widths[i])
    cell_text(cell,label,True)
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EAF0F5');cell._tc.get_or_add_tcPr().append(shd)
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

for name, colors, palette, impression, home, resources in options:
    row=t.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for i, cell in enumerate(row.cells):
        cell.width=Inches(widths[i])
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=cell_text(row.cells[0],name,True)
    p.add_run('\n')
    for color in colors:
        r=p.add_run('■ ');r.font.color.rgb=RGBColor.from_string(color if color!='FFFFFF' else 'C4CCD3');r.font.size=Pt(13)
    r=p.add_run('\n'+palette);r.font.size=Pt(9);r.bold=False
    for cell,text in zip(row.cells[1:],[impression,home,resources]):cell_text(cell,text)

p=doc.add_paragraph('工作量为相对比较，最终取决于学长模板及学校站群支持情况。色块仅为方向示意。')
p.paragraph_format.space_before=Pt(5)
for r in p.runs:r.font.size=Pt(9)

doc.add_heading('建议选择',level=2)
p=doc.add_paragraph()
p.add_run('建议以 B 现代科技风为主，采用 C 的清晰排版。').bold=True
p.add_run('首版可用静态芯片主题图或实验室照片，配少量切换效果；完整电路板动画留作后续可选项，先满足15号演示需要。')

doc.add_heading('老师选择与意见',level=2)
for text in [
    '主风格：□ A 稳重学院风   □ B 现代科技风   □ C 简洁学术风   □ D 大图展示风',
    '首页配图：□ 校园或学院实景   □ 实验室实景   □ 芯片主题图',
    '动效偏好：□ 基本无动效   □ 少量切换效果   □ 希望突出动画或视频',
    '补充意见：________________________________________________________',
    '确认人：________________    日期：________________',
]:doc.add_paragraph(text)

doc.core_properties.subject = '三位教师资料人工核验及学院官网设计风格选择'
doc.core_properties.title = '三位教师简历与照片及网站风格选型'
doc.save(output)

saved=Document(output)
saved_blocks=[etree.tostring(e) for e in saved._element.body if e.tag!=qn('w:sectPr')]
assert saved_blocks[:len(original_blocks)]==original_blocks, 'Original body changed'
with ZipFile(source) as old, ZipFile(output) as new:
    photos=[n for n in old.namelist() if n.startswith('word/media/')]
    assert len(photos)==3
    for name in photos:assert old.read(name)==new.read(name)
assert all(item[0] in '\n'.join(c.text for row in saved.tables[-1].rows for c in row.cells) for item in options)
print(output)
print(f'Preserved all {len(original_blocks)} existing body blocks and 3 original photos; added 4 style options')
