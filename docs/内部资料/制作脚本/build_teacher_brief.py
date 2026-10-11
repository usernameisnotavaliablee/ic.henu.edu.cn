"""Prepare a short, teacher-facing confirmation packet with visual style choices."""
from pathlib import Path
import sys
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
docs = next(p for p in Path(__file__).resolve().parents if p.name=='docs')
output = docs/'发给老师'/'老师确认稿_教师资料与网站风格.docx'
asset_dir=docs/'assets'/'design'
asset_dir.mkdir(parents=True,exist_ok=True)
board_path=asset_dir/'网站风格方向示意.png'

# These are original layout sketches, not screenshots of completed websites.
board=Image.new('RGB',(1600,1100),'white')
draw=ImageDraw.Draw(board)
font_path='C:/Windows/Fonts/msyh.ttc'
bold_path='C:/Windows/Fonts/msyhbd.ttc'
def font(n,bold=False):return ImageFont.truetype(bold_path if bold else font_path,n)
def txt(x,y,t,n=24,fill='#23313d',bold=False):draw.text((x,y),t,font=font(n,bold),fill=fill)
styles=[
    ('A  稳重学院风','红白为主  规整正式','#923642','#F7F3F0'),
    ('B  蓝白科技风','蓝白为主  突出芯片特色','#245B97','#ECF3F9'),
    ('C  简洁学术风','白底为主  清楚好读','#315779','#F3F5F7'),
    ('D  大图展示风','大幅照片  形象突出','#163B4B','#F0F4F6'),
]
for i,(name,desc,accent,pale) in enumerate(styles):
    ox=(i%2)*820;oy=(i//2)*550
    txt(ox+8,oy+6,name,38,'#000000',True)
    txt(ox+8,oy+58,desc,29)
    x=ox+8;y=oy+108;w=760;h=404
    draw.rectangle((x,y,x+w,y+h),fill='white',outline='#CCD3DA',width=2)
    if i==0:
        draw.rectangle((x+2,y+2,x+w-2,y+54),fill=accent)
        txt(x+20,y+10,'学院名称',26,'white',True)
        txt(x+20,y+66,'学院概况   师资队伍   教学科研   新闻通知',20)
        draw.rectangle((x+18,y+102,x+w-18,y+254),fill=pale)
        txt(x+260,y+158,'学院实景照片',29,accent)
        txt(x+20,y+273,'学院新闻',25,accent,True)
        txt(x+415,y+273,'通知公告',25,accent,True)
        for yy in [323,351,379]:
            draw.line((x+20,y+yy,x+351,y+yy),fill='#D6DCE2',width=6)
            draw.line((x+417,y+yy,x+729,y+yy),fill='#D6DCE2',width=6)
    elif i==1:
        txt(x+20,y+13,'学院名称',26,accent,True)
        draw.rectangle((x+18,y+63,x+w-18,y+245),fill=accent)
        txt(x+46,y+112,'集成电路主题',31,'white',True)
        txt(x+46,y+160,'简短学院介绍',22,'#DEEBF7')
        cx=x+577;cy=y+155
        draw.rectangle((cx-44,cy-44,cx+44,cy+44),fill='#DCEEFF',outline='white',width=3)
        for d in [-28,-10,10,28]:
            draw.line((cx+d,cy-70,cx+d,cy-44),fill='#A4D4FF',width=4)
            draw.line((cx+d,cy+44,cx+d,cy+70),fill='#A4D4FF',width=4)
            draw.line((cx-70,cy+d,cx-44,cy+d),fill='#A4D4FF',width=4)
            draw.line((cx+44,cy+d,cx+70,cy+d),fill='#A4D4FF',width=4)
        for j,label in enumerate(['研究方向','师资队伍','学院动态']):
            xx=x+18+j*243
            draw.rectangle((xx,y+272,xx+225,y+382),fill=pale)
            txt(xx+54,y+309,label,26,accent,True)
    elif i==2:
        txt(x+28,y+20,'学院名称',27,accent,True)
        draw.line((x+28,y+70,x+w-28,y+70),fill='#D6DCE2',width=2)
        txt(x+30,y+99,'学院简介',30,'#23313D',True)
        for yy,ww in [(156,450),(186,630),(216,570)]:
            draw.line((x+30,y+yy,x+30+ww,y+yy),fill='#DDE2E7',width=7)
        txt(x+30,y+267,'教师与研究方向',28,accent,True)
        for yy in [322,359]:draw.line((x+30,y+yy,x+685,y+yy),fill='#E3E7EB',width=7)
    else:
        draw.rectangle((x+2,y+2,x+w-2,y+298),fill=accent)
        txt(x+24,y+20,'学院名称',27,'white',True)
        txt(x+238,y+121,'学院或实验室大图',31,'white',True)
        txt(x+282,y+178,'配一句学院介绍',23,'#D8E5EB')
        txt(x+28,y+322,'学院概况',25,accent,True)
        txt(x+282,y+322,'教学科研',25,accent,True)
        txt(x+541,y+322,'新闻通知',25,accent,True)
        draw.line((x+28,y+372,x+w-28,y+372),fill='#DDE3E7',width=7)
board.save(board_path)

doc=Document(docs/'历史版本'/'三位教师简历与照片_人工核验稿.docx')
for child in list(doc._element.body):
    if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
doc.core_properties.title='教师资料与网站风格确认稿'
doc.core_properties.subject='供学院老师核对三位教师资料并选择网站风格'
header=doc.sections[0].header.paragraphs[0]
header.text='学院网站建设    请老师核对与选择'
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.18
normal.space_after=Pt(6)

def p(text,bold=False):
    para=doc.add_paragraph()
    r=para.add_run(text);r.bold=bold
    return para
def h(text):return doc.add_heading(text,level=2)
def link(para,label,url):
    rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    node=OxmlElement('w:hyperlink');node.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'24527A');pr.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);node.append(r);para._p.append(node)

people=[
    dict(name='张锦龙',photo='张锦龙.png',title='教授、博士生导师',
         direction='光纤传感与光子芯片',contact='邮箱 zjl@henu.edu.cn；电话 135-6950-9930',
         bullets=[
             '北京邮电大学电磁场与微波技术专业博士。此前在长沙理工大学取得学士学位，在河南大学取得硕士学位。',
             '现有资料列有河南省特聘教授、河南省高层次人才、河南省高校科技创新团队带头人等称号。',
             '承担《电路原理》《电子测量》《现代交换原理》《新型光传感技术及其应用》等课程。',
             '现有资料记载：发表论文40余篇，授权国家发明专利16件，主持国家基金面上项目2项。',
         ],source='https://phye.henu.edu.cn/info/1159/5480.htm',photo_note='现有照片较小，如有高清原图，请一并提供。'),
    dict(name='郑海务',photo='郑海务.jpg',title='二级教授、博士生导师',
         direction='微纳传感、能量收集器件、半导体器件及光电探测等',
         contact='邮箱 zhenghw@henu.edu.cn；hwxyzheng@163.com',
         bullets=[
             '现有资料列为河南大学集成电路科学与工程省重点学科带头人、集成电路设计与集成系统专业负责人等，请核对现任职务。',
             '现有资料列有中原基础研究领军人才、河南省特聘教授、河南省杰出青年基金获得者等称号。',
             '承担《薄膜物理》《集成电路工艺原理》《科技论文写作》等课程。',
             '现有资料记载：主持国家自然科学基金项目5项，第一发明人获授权国家发明专利17件。',
         ],source='https://phye.henu.edu.cn/info/1159/5688.htm',photo_note='现有照片清晰度有限，如有高清原图，请一并提供。'),
    dict(name='程纲',photo='程纲.jpg',title='教授、博士生导师',
         direction='请老师补充或确认用于学院官网的一句研究方向介绍。',
         contact='邮箱 chenggang@henu.edu.cn',
         bullets=[
             '2000年、2003年分别在河南大学取得学士和硕士学位，2008年在吉林大学取得凝聚态物理专业博士学位。',
             '现有资料所列单位为河南大学特种功能材料教育部重点实验室，请核对现任职务及对外展示的单位名称。',
             '曾在香港中文大学任研究助理，在佐治亚理工学院任访问学者。',
             '现有资料列有国家优秀青年基金获得者（2015年）、河南省科技创新人才等。主讲课程和招生介绍可另行补充。',
         ],source='https://lab.henu.edu.cn/zw/ktz/c_g_Cheng_G/ktzzjj.htm',photo_note='已有高清原图，请确认这张照片是否可以使用。'),
]

for i,person in enumerate(people):
    title=doc.add_paragraph(person['name']+'  教师资料确认',style='Title')
    if i:title.paragraph_format.page_break_before=True
    p('以下是根据已有公开资料整理的简介，请核对最新情况后用于学院官网。')
    pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pic=pp.add_run().add_picture(str(docs/'assets'/'teachers'/person['photo']),height=Inches(1.2))
    pic._inline.docPr.set('descr',person['name']+'现有照片')
    pp.paragraph_format.keep_with_next=True
    note=p(person['photo_note']);note.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in note.runs:r.font.size=Pt(10)
    for label,key in [('职称','title'),('研究方向','direction'),('联系方式','contact')]:
        para=doc.add_paragraph();para.add_run(label+'：').bold=True;para.add_run(person[key])
    h('简历要点')
    for text in person['bullets']:doc.add_paragraph(text,style='List Bullet')
    h('请老师确认')
    for text in [
        '是否在学院网站展示这位老师：□ 是    □ 否',
        '上述资料：□ 可以使用    □ 需要修改（请直接改在文中）',
        '照片：□ 可以使用    □ 将提供新照片',
        '联系方式：□ 上述信息均可公开    □ 请仅公开：________________',
        '其他意见：______________________________________________________',
    ]:p(text)
    foot=doc.add_paragraph();foot.add_run('资料来源：');link(foot,'点击查看原教师介绍页面',person['source'])
    for r in foot.runs:r.font.size=Pt(9)

title=doc.add_paragraph('网站风格  请选一个喜欢的方向',style='Title')
title.paragraph_format.page_break_before=True
p('下面是四种大致效果示意。先选整体感觉，具体图片和文字之后再调整。')
picp=doc.add_paragraph();picp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pic=picp.add_run().add_picture(str(board_path),width=Inches(6.8))
pic._inline.docPr.set('descr','四种网站风格的布局示意，分别为稳重学院风、蓝白科技风、简洁学术风和大图展示风')
p('建议选择 B 蓝白科技风：比较贴合集成电路学院的特点，也能保持页面清楚好读。')
h('老师的选择')
for text in [
    '我更喜欢：□ A 稳重学院风    □ B 蓝白科技风',
    '                □ C 简洁学术风    □ D 大图展示风',
    '                □ E 其他风格：______________________________________',
    '参考网站或喜欢的效果：____________________________________________',
    '希望调整的地方：________________________________________________',
    '确认人：________________    日期：________________',
]:p(text)
doc.save(output)
check=Document(output)
assert len(check.inline_shapes)==4
assert len([para for para in check.paragraphs if para.style.name=='Title'])==4
assert all(n in '\n'.join(p.text for p in check.paragraphs) for n in ['张锦龙','郑海务','程纲','老师的选择'])
print(output)
print('4 intended pages; 3 faculty photos and 1 style illustration embedded')
