"""Add the requested open-ended choice while preserving teacher edits."""
from pathlib import Path
import sys
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')
docs=next(p for p in Path(__file__).resolve().parents if p.name=='docs')
path=docs/'发给老师'/'老师确认稿_教师资料与网站风格.docx'
if not path.exists():path=docs/'老师确认稿_教师资料与网站风格.docx'
doc=Document(path)
old_blocks=[etree.tostring(x) for x in doc._element.body]
anchor=next(p for p in doc.paragraphs if '□ C 简洁学术风' in p.text and '□ D 大图展示风' in p.text)
labels=['                □ E 其他风格：______________________________________',
        '参考网站或喜欢的效果：____________________________________________']
if not any('□ E 其他风格' in p.text for p in doc.paragraphs):
    for text in labels:
        element=OxmlElement('w:p')
        anchor._p.addnext(element)
        new=Paragraph(element,anchor._parent)
        new.style=anchor.style
        new.add_run(text)
        anchor=new
doc.save(path)
saved=Document(path)
new_blocks=[etree.tostring(x) for x in saved._element.body]
# Every original body block must still be present, in original order.
position=0
for old in old_blocks:
    while position<len(new_blocks) and new_blocks[position]!=old:position+=1
    assert position<len(new_blocks),'An original body block changed'
    position+=1
assert len(saved.inline_shapes)==4
assert '□ E 其他风格' in '\n'.join(p.text for p in saved.paragraphs)
print(f'Updated {path}; original body and four images preserved')
